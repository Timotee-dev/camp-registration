import io
import json
from datetime import timedelta

from django.contrib.auth.models import User
from django.core.cache import cache
from django.test import TestCase, override_settings
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from .answers import normalise_phone
from .models import Event, Question, Submission, UniqueAnswer


def make_user(email="ada@example.com", name="Ada Obi", pw="Str0ng-pass!"):
    return User.objects.create_user(username=email, email=email, password=pw, first_name=name)


class AccountTests(TestCase):
    def test_signup_logs_in_and_lands_on_dashboard(self):
        r = self.client.post(reverse("accounts:signup"),
                             {"full_name": "  Tolu   Adeyemi ", "email": "Tolu@Example.com", "password": "Camp-Meeting-26"})
        self.assertRedirects(r, reverse("events:dashboard"))
        u = User.objects.get()
        self.assertEqual((u.first_name, u.email, u.username), ("Tolu Adeyemi", "tolu@example.com", "tolu@example.com"))
        self.assertContains(self.client.get(reverse("events:dashboard")), "Create your first event")

    def test_duplicate_email_and_weak_password_rejected(self):
        make_user("tolu@example.com")
        r = self.client.post(reverse("accounts:signup"), {"full_name": "X Y", "email": "TOLU@example.com", "password": "Camp-Meeting-26"})
        self.assertContains(r, "already exists")
        r = self.client.post(reverse("accounts:signup"), {"full_name": "X Y", "email": "new@example.com", "password": "12345678"})
        self.assertContains(r, "too common")
        self.assertEqual(User.objects.count(), 1)

    def test_login_with_email_case_insensitive(self):
        make_user()
        r = self.client.post(reverse("accounts:login"), {"email": "ADA@example.com", "password": "Str0ng-pass!"})
        self.assertRedirects(r, reverse("events:dashboard"))
        r = self.client_class().post(reverse("accounts:login"), {"email": "ada@example.com", "password": "wrong"})
        self.assertContains(r, "Incorrect email or password")

    def test_login_next_is_safe(self):
        make_user()
        r = self.client.post(reverse("accounts:login") + "?next=https://evil.com",
                             {"email": "ada@example.com", "password": "Str0ng-pass!", "next": "https://evil.com"})
        self.assertRedirects(r, reverse("events:dashboard"))

    def test_dashboard_requires_login(self):
        r = self.client.get(reverse("events:dashboard"))
        self.assertEqual(r.status_code, 302)
        self.assertIn(reverse("accounts:login"), r.url)


class BuilderTests(TestCase):
    def setUp(self):
        self.user = make_user()
        self.client.force_login(self.user)

    def create_event(self, **extra):
        data = {"title": "Youth Camp 2026", "description": "Come and see", "venue": "Ede", "accent": "#1f4e3d",
                "organiser_name": "RCCG", "contact": "0803", "starts_at": "", "ends_at": ""}
        data.update(extra)
        r = self.client.post(reverse("events:create"), data)
        return r, Event.objects.order_by("-id").first()

    def test_create_event_seeds_starter_questions(self):
        r, ev = self.create_event()
        self.assertRedirects(r, reverse("events:build", args=[ev.pk]))
        self.assertEqual([q.label for q in ev.questions.all()], ["Full name", "Phone number", "Email address", "Home address"])
        self.assertEqual(len(ev.code), 8)

    def test_end_before_start_rejected(self):
        r, ev = self.create_event(starts_at="2026-12-20T10:00", ends_at="2026-12-19T10:00")
        self.assertContains(r, "can&#x27;t be before the start")
        self.assertIsNone(ev)

    def put(self, ev, questions, client=None):
        return (client or self.client).put(reverse("events:api_questions", args=[ev.pk]),
                                           data=json.dumps({"questions": questions}), content_type="application/json")

    def test_save_questions_create_update_reorder_delete(self):
        _, ev = self.create_event()
        qs = list(ev.questions.all())
        payload = [
            {"id": qs[1].id, "qtype": "phone", "label": "Phone", "required": True, "unique": True},
            {"id": qs[0].id, "qtype": "short_text", "label": "Name", "required": True},
            {"qtype": "dropdown", "label": "  Gender ", "options": ["Male", " Female ", ""], "required": True},
        ]
        r = self.put(ev, payload)
        self.assertEqual(r.status_code, 200, r.content)
        out = r.json()["questions"]
        self.assertEqual([q["label"] for q in out], ["Phone", "Name", "Gender"])
        self.assertEqual(out[2]["options"], ["Male", "Female"])
        self.assertTrue(out[0]["unique"])
        self.assertFalse(Question.objects.filter(pk__in=[qs[2].id, qs[3].id]).exists())  # unanswered -> deleted

    def test_invalid_questions_rejected(self):
        _, ev = self.create_event()
        r = self.put(ev, [{"qtype": "single_choice", "label": "Pick", "options": ["A", "a"]},
                          {"qtype": "short_text", "label": "   "}])
        self.assertEqual(r.status_code, 400)
        errs = r.json()["errors"]
        self.assertEqual([e["index"] for e in errs], [0, 1])
        self.assertEqual(ev.questions.count(), 4)  # unchanged

    def test_unique_only_for_allowed_types(self):
        _, ev = self.create_event()
        r = self.put(ev, [{"qtype": "long_text", "label": "Address", "unique": True}])
        self.assertFalse(r.json()["questions"][0]["unique"])

    def test_removing_answered_question_archives_it(self):
        _, ev = self.create_event()
        q = ev.questions.first()
        Submission.objects.create(event=ev, number=1, answers={str(q.id): "Ada"})
        others = [{"id": x.id, "qtype": x.qtype, "label": x.label} for x in ev.questions.exclude(pk=q.pk)]
        self.assertEqual(self.put(ev, others).status_code, 200)
        q.refresh_from_db()
        self.assertTrue(q.is_archived)
        r = self.client.get(reverse("events:responses", args=[ev.pk]))
        self.assertContains(r, "(removed)")

    def test_other_users_cannot_touch_event(self):
        _, ev = self.create_event()
        other = self.client_class()
        other.force_login(make_user("bad@example.com"))
        for name in ["build", "share", "settings", "responses", "summary"]:
            self.assertEqual(other.get(reverse(f"events:{name}", args=[ev.pk])).status_code, 404, name)
        self.assertEqual(other.get(reverse("events:export", args=[ev.pk, "xlsx"])).status_code, 404)
        self.assertEqual(self.put(ev, [], client=other).status_code, 404)
        self.assertEqual(self.client_class().put(reverse("events:api_questions", args=[ev.pk]), "{}",
                                                 content_type="application/json").status_code, 403)

    def test_share_page_has_link_and_qr(self):
        _, ev = self.create_event()
        r = self.client.get(reverse("events:share", args=[ev.pk]))
        self.assertContains(r, f"/r/{ev.code}/")
        self.assertContains(r, "<svg")
        png = self.client.get(reverse("events:qr_png", args=[ev.pk]))
        self.assertEqual(png["Content-Type"], "image/png")

    def test_duplicate_and_delete(self):
        _, ev = self.create_event()
        self.client.post(reverse("events:duplicate", args=[ev.pk]))
        copy = Event.objects.exclude(pk=ev.pk).get()
        self.assertEqual(copy.questions.count(), 4)
        self.assertNotEqual(copy.code, ev.code)
        self.client.post(reverse("events:delete", args=[ev.pk]), {"confirm": "nope"})
        self.assertTrue(Event.objects.filter(pk=ev.pk).exists())
        self.client.post(reverse("events:delete", args=[ev.pk]), {"confirm": "DELETE"})
        self.assertFalse(Event.objects.filter(pk=ev.pk).exists())

    def test_settings_rules(self):
        _, ev = self.create_event()
        r = self.client.post(reverse("events:settings", args=[ev.pk]),
                             {"section": "rules", "r-max_responses": "50", "r-success_message": "See you!"})
        self.assertEqual(r.status_code, 302)
        ev.refresh_from_db()
        self.assertEqual((ev.is_open, ev.max_responses, ev.success_message), (False, 50, "See you!"))


@override_settings(SUBMISSION_RATE_LIMIT=1000)
class PublicFormTests(TestCase):
    def setUp(self):
        cache.clear()
        self.owner = make_user()
        self.ev = Event.objects.create(owner=self.owner, title="Camp", description="Hello", success_message="Bring a Bible")
        mk = lambda **k: Question.objects.create(event=self.ev, **k)
        self.name = mk(order=0, qtype="short_text", label="Full name")
        self.phone = mk(order=1, qtype="phone", label="Phone", unique=True)
        self.email = mk(order=2, qtype="email", label="Email", required=False)
        self.gender = mk(order=3, qtype="single_choice", label="Gender", options=["Male", "Female"])
        self.diet = mk(order=4, qtype="multi_choice", label="Diet", options=["Rice", "Beans", "Yam"], required=False)
        self.age = mk(order=5, qtype="number", label="Age", required=False)
        self.dob = mk(order=6, qtype="date", label="Arrival", required=False)
        self.url = reverse("public:form", args=[self.ev.code])

    def data(self, **over):
        d = {f"q_{self.name.id}": "  Tolu   Adeyemi ", f"q_{self.phone.id}": "+234 803 123 4567",
             f"q_{self.email.id}": "TOLU@Example.com", f"q_{self.gender.id}": "Female",
             f"q_{self.diet.id}": ["Yam", "Rice"], f"q_{self.age.id}": "1,200", f"q_{self.dob.id}": "2026-12-20"}
        d.update(over)
        return d

    def test_intro_and_steps_render(self):
        r = self.client.get(self.url)
        self.assertContains(r, "Start registration")
        self.assertContains(r, "Hello")
        self.assertContains(r, 'data-step="7"')
        self.assertContains(r, "Check your answers")

    def test_successful_submission(self):
        r = self.client.post(self.url, self.data())
        self.assertRedirects(r, reverse("public:thanks", args=[self.ev.code]))
        s = Submission.objects.get()
        self.assertEqual(s.number, 1)
        self.assertEqual(s.answers, {
            str(self.name.id): "Tolu Adeyemi", str(self.phone.id): "08031234567", str(self.email.id): "tolu@example.com",
            str(self.gender.id): "Female", str(self.diet.id): ["Rice", "Yam"], str(self.age.id): "1200",
            str(self.dob.id): "2026-12-20",
        })
        t = self.client.get(r.url)
        self.assertContains(t, "#1")
        self.assertContains(t, "Bring a Bible")

    def test_optional_blank_answers_omitted(self):
        d = self.data(**{f"q_{self.email.id}": "", f"q_{self.age.id}": "", f"q_{self.dob.id}": ""})
        d.pop(f"q_{self.diet.id}")
        self.client.post(self.url, d)
        self.assertEqual(set(Submission.objects.get().answers), {str(self.name.id), str(self.phone.id), str(self.gender.id)})

    def test_validation_errors_kept_and_shown(self):
        r = self.client.post(self.url, self.data(**{
            f"q_{self.name.id}": "", f"q_{self.phone.id}": "abc", f"q_{self.email.id}": "nope",
            f"q_{self.gender.id}": "Other", f"q_{self.diet.id}": ["Pizza"], f"q_{self.age.id}": "ten",
            f"q_{self.dob.id}": "2026-02-31"}))
        self.assertEqual(r.status_code, 200)
        for msg in ["This question is required", "valid phone number", "valid email", "choose from the options",
                    "Enter a number", "valid date"]:
            self.assertContains(r, msg)
        self.assertContains(r, 'data-has-errors="1"')
        self.assertFalse(Submission.objects.exists())

    def test_duplicate_phone_any_format(self):
        self.client.post(self.url, self.data())
        r = self.client.post(self.url, self.data(**{f"q_{self.phone.id}": "0803-123-4567"}))
        self.assertContains(r, "already been used")
        self.assertEqual(Submission.objects.count(), 1)

    def test_numbers_increment_and_survive_deletes(self):
        self.client.post(self.url, self.data())
        self.client.post(self.url, self.data(**{f"q_{self.phone.id}": "07011112222"}))
        Submission.objects.get(number=2).delete()
        self.client.post(self.url, self.data(**{f"q_{self.phone.id}": "09011112222"}))
        self.assertEqual(sorted(Submission.objects.values_list("number", flat=True)), [1, 3])

    def test_deleting_registration_frees_unique_answer(self):
        self.client.post(self.url, self.data())
        Submission.objects.get().delete()
        self.assertFalse(UniqueAnswer.objects.exists())
        self.client.post(self.url, self.data())
        self.assertEqual(Submission.objects.count(), 1)

    def test_closed_full_ended(self):
        self.ev.is_open = False
        self.ev.save()
        self.assertContains(self.client.get(self.url), "Registration is closed")
        self.client.post(self.url, self.data())
        self.assertFalse(Submission.objects.exists())

        self.ev.is_open, self.ev.max_responses = True, 1
        self.ev.save()
        self.client.post(self.url, self.data())
        self.assertContains(self.client.get(self.url), "Registration is full")

        self.ev.max_responses, self.ev.close_at = None, timezone.now() - timedelta(minutes=1)
        self.ev.save()
        self.assertContains(self.client.get(self.url), "Registration has ended")

    def test_empty_form_not_accepting(self):
        ev = Event.objects.create(owner=self.owner, title="Empty")
        self.assertContains(self.client.get(reverse("public:form", args=[ev.code])), "isn't ready yet")

    def test_honeypot_drops_silently(self):
        r = self.client.post(self.url, self.data(website_url="http://spam"))
        self.assertEqual(r.status_code, 302)
        self.assertFalse(Submission.objects.exists())

    @override_settings(SUBMISSION_RATE_LIMIT=2)
    def test_rate_limit(self):
        for p in ["07011110001", "07011110002", "07011110003"]:
            r = self.client.post(self.url, self.data(**{f"q_{self.phone.id}": p}))
        self.assertContains(r, "Too many registrations")
        self.assertEqual(Submission.objects.count(), 2)

    def test_unknown_code_404(self):
        self.assertEqual(self.client.get(reverse("public:form", args=["zzzzzzzz"])).status_code, 404)

    def test_turning_unique_on_later_covers_existing(self):
        self.phone.unique = False
        self.phone.save()
        UniqueAnswer.objects.all().delete()
        self.client.post(self.url, self.data())
        self.client.force_login(self.owner)
        payload = [{"id": q.id, "qtype": q.qtype, "label": q.label, "options": q.options, "required": q.required,
                    "unique": q.id == self.phone.id} for q in self.ev.questions.all()]
        self.client.put(reverse("events:api_questions", args=[self.ev.pk]), json.dumps({"questions": payload}),
                        content_type="application/json")
        r = self.client_class().post(self.url, self.data())
        self.assertContains(r, "already been used")


class ExportAndViewTests(TestCase):
    def setUp(self):
        self.owner = make_user()
        self.client.force_login(self.owner)
        self.ev = Event.objects.create(owner=self.owner, title="Camp 2026", accent="#2b4acb")
        self.n = Question.objects.create(event=self.ev, order=0, qtype="short_text", label="Name")
        self.p = Question.objects.create(event=self.ev, order=1, qtype="phone", label="Phone")
        self.g = Question.objects.create(event=self.ev, order=2, qtype="yes_no", label="First time?")
        self.a = Question.objects.create(event=self.ev, order=3, qtype="number", label="Age", required=False)
        self.old = Question.objects.create(event=self.ev, order=4, qtype="short_text", label="Church", is_archived=True)
        for i, (name, ft) in enumerate([("Ada", "Yes"), ("Bayo", "No"), ("Chika", "Yes")], start=1):
            Submission.objects.create(event=self.ev, number=i, answers={
                str(self.n.id): name, str(self.p.id): f"0803000000{i}", str(self.g.id): ft, str(self.a.id): str(20 + i),
                str(self.old.id): "RCCG"})

    def test_xlsx_export(self):
        r = self.client.get(reverse("events:export", args=[self.ev.pk, "xlsx"]))
        wb = load_workbook(io.BytesIO(r.content))
        ws = wb["Registrations"]
        self.assertEqual([c.value for c in ws[4]], ["No.", "Registered at", "Name", "Phone", "First time?", "Age", "Church (removed)"])
        self.assertEqual([c.value for c in ws[5]][2:], ["Ada", "08030000001", "Yes", 21, "RCCG"])
        self.assertEqual(ws.cell(5, 4).number_format, "@")
        self.assertEqual(ws.max_row, 7)
        sm = [r[0] for r in wb["Summary"].iter_rows(values_only=True) if r]
        self.assertIn("First time?", sm)

    def test_csv_export(self):
        r = self.client.get(reverse("events:export", args=[self.ev.pk, "csv"]))
        lines = r.content.decode("utf-8-sig").strip().splitlines()
        self.assertEqual(len(lines), 4)
        self.assertTrue(lines[1].startswith("1,"))
        self.assertIn("08030000001", lines[1])

    def test_responses_search_detail_delete(self):
        r = self.client.get(reverse("events:responses", args=[self.ev.pk]), {"q": "bayo"})
        self.assertEqual(r.context["filtered"], 1)
        r = self.client.get(reverse("events:responses", args=[self.ev.pk]), {"q": "08030000003"})
        self.assertEqual(r.context["filtered"], 1)
        r = self.client.get(reverse("events:responses", args=[self.ev.pk]), {"q": "#2"})
        self.assertEqual(r.context["filtered"], 1)
        d = self.client.get(reverse("events:response", args=[self.ev.pk, 2]))
        self.assertContains(d, "Bayo")
        self.client.post(reverse("events:response_delete", args=[self.ev.pk, 2]))
        self.assertEqual(self.ev.submissions.count(), 2)

    def test_summary(self):
        r = self.client.get(reverse("events:summary", args=[self.ev.pk]))
        tallies = {q.label: rows for q, rows, _ in r.context["tallies"]}
        self.assertEqual(tallies["First time?"], [("Yes", 2, 67), ("No", 1, 33)])

    def test_toggle(self):
        self.client.post(reverse("events:toggle", args=[self.ev.pk]))
        self.ev.refresh_from_db()
        self.assertFalse(self.ev.is_open)


class PhoneTests(TestCase):
    def test_normalise(self):
        for raw in ["08031234567", "0803 123 4567", "+2348031234567", "234 803 123 4567", "8031234567", "(0803) 123-4567"]:
            self.assertEqual(normalise_phone(raw), "08031234567", raw)
        self.assertEqual(normalise_phone("+44 20 7946 0958"), "+442079460958")


class HealthCheckTests(TestCase):
    def test_healthz_ignores_host_and_https_checks(self):
        r = self.client.get("/healthz", HTTP_HOST="10.0.0.5:10000")
        self.assertEqual((r.status_code, r.content), (200, b"ok"))


class ErrorPageTests(TestCase):
    def test_db_health_and_404(self):
        r = self.client.get("/healthz?db=1")
        self.assertIn(b"db: ok", r.content)
        r = self.client.get("/r/doesnotexist/")
        self.assertEqual(r.status_code, 404)
